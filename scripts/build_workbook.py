"""Builds the Google Sheets workbook (as .xlsx) for NBA and NFL game log
lookups. Run with:  python build_workbook.py test   (data filled in, for checking)
                    python build_workbook.py template (IMPORTDATA links, what you use)
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
from schema import NBA_COLS, NBA_STATS, NFL_COLS, NFL_STATS  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
MAXR = 60000
INPUT = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor="D9E2F3")
BOLD = Font(bold=True)
NBA_TEAMS = "ATL BOS BKN CHA CHI CLE DAL DEN DET GS HOU IND LAC LAL MEM MIA MIL MIN NO NY OKC ORL PHI PHX POR SAC SA TOR UTAH WSH".split()
NFL_TEAMS = "ARI ATL BAL BUF CAR CHI CIN CLE DAL DEN DET GB HOU IND JAX KC LA LAC LV MIA MIN NE NO NYG NYJ PHI PIT SEA SF TB TEN WAS".split()

SPORTS = {
    "NBA": dict(cols=NBA_COLS, stats=NBA_STATS, teams=NBA_TEAMS, usage="MIN",
                usage_label="Avg MIN", cur="Settings!$B$10", first=5, long=[5, 10, 20],
                default=("Scottie Barnes", "PTS", 20, 1.8, 25),
                splits=[("Home", [("home", '1')]), ("Away", [("home", '0')]),
                        ("vs opponent", "OPP"),
                        ("Starter", [("starter", '1')]), ("Bench", [("starter", '0')]),
                        ("Close games (margin 9 or less)", [("abs_margin", '"<=9"')]),
                        ("Blowouts (margin 15 or more)", [("abs_margin", '">=15"')]),
                        ("Back to back", [("b2b", '1')]),
                        ("Played at least X min", "USAGE")]),
    "NFL": dict(cols=NFL_COLS, stats=NFL_STATS, teams=NFL_TEAMS, usage="OPP",
                usage_label="Avg OPP", cur="Settings!$B$11", first=3, long=[3, 5, 8],
                default=("Patrick Mahomes", "PASS_YDS", 250, 1.8, 0),
                splits=[("Home", [("home", '1')]), ("Away", [("home", '0')]),
                        ("vs opponent", "OPP"),
                        ("Team favored", [("fav_by", '">0"')]),
                        ("Team underdog", [("fav_by", '"<0"')]),
                        ("Close games (margin 7 or less)", [("abs_margin", '"<=7"')]),
                        ("Blowouts (margin 14 or more)", [("abs_margin", '">=14"')]),
                        ("Had at least X opportunities", "USAGE")]),
}


def hdr(ws, row, labels, widths=None):
    for i, t in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = BOLD
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")


def build(mode, out):
    wb = Workbook()
    wb.remove(wb.active)
    readme = wb.create_sheet("README")
    settings = wb.create_sheet("Settings")
    for sp in ("NBA", "NFL"):
        wb.create_sheet(f"{sp}_Dash")
        wb.create_sheet(f"{sp}_Ticket")
    for sp in ("NBA", "NFL"):
        wb.create_sheet(f"Data_{sp}")
        wb.create_sheet(f"Players_{sp}")

    # ---------- README ----------
    lines = [
        "Game log lookup (NBA and NFL)",
        "",
        "NBA_Dash / NFL_Dash: pick a player, a stat and a line. Yellow cells are inputs.",
        "NBA_Ticket / NFL_Ticket: list your legs, see hit rates, break even and a rough price check.",
        "Data_ and Players_ tabs fill themselves from GitHub every night. Do not type in them.",
        "Settings: the four data links (set once), the haircut and the small sample size.",
        "",
        "How to read it",
        "Line means X or more. For Over 4.5 enter 5. For Under, use 1 minus the hit rate.",
        "Break even is 1 divided by the decimal price. Edge is hit rate minus break even.",
        "Hit rate is history, not a forecast. Small samples (under 8 games) are flagged.",
        "Splits default to this season plus last season. Change the 'Splits use seasons from' cell.",
        "NFL caveat: a game where a player has no recorded stats is missing, so very low lines look too easy.",
        "Ticket tab multiplies leg odds as if independent. Same game legs are correlated, so treat it as a rough guide.",
        "Games with no box score yet or an ESPN outage are simply missing until the next nightly run.",
    ]
    for i, t in enumerate(lines, 1):
        readme.cell(row=i, column=1, value=t)
    readme["A1"].font = Font(bold=True, size=14)
    readme.column_dimensions["A"].width = 110

    # ---------- Settings ----------
    settings["A1"] = "Settings"
    settings["A1"].font = Font(bold=True, size=14)
    settings["A2"] = "Paste the four raw.githubusercontent.com links (see GUIDE). Set once."
    names = ["NBA game log link", "NBA players link", "NFL game log link", "NFL players link"]
    files = ["nba_gamelog.csv", "nba_players.csv", "nfl_gamelog.csv", "nfl_players.csv"]
    for i, (n, f) in enumerate(zip(names, files)):
        settings.cell(row=3 + i, column=1, value=n)
        c = settings.cell(row=3 + i, column=2, value=f"https://raw.githubusercontent.com/YOURNAME/YOURREPO/main/data/{f}")
        c.fill = INPUT
    settings["A8"] = "Haircut on auto estimate (relative)"
    settings["B8"] = 0.10
    settings["B8"].fill = INPUT
    settings["B8"].number_format = "0%"
    settings["A9"] = "Small sample below (games)"
    settings["B9"] = 8
    settings["B9"].fill = INPUT
    settings["A10"] = "Current NBA season (end year)"
    settings["B10"] = f"=MAX(Data_NBA!$B$2:$B${MAXR})"
    settings["A11"] = "Current NFL season"
    settings["B11"] = f"=MAX(Data_NFL!$B$2:$B${MAXR})"
    settings.column_dimensions["A"].width = 36
    settings.column_dimensions["B"].width = 90
    # team lists for dropdowns
    for j, sp in enumerate(("NBA", "NFL")):
        settings.cell(row=14, column=1 + j, value=f"{sp} teams").font = BOLD
        for i, t in enumerate(SPORTS[sp]["teams"]):
            settings.cell(row=15 + i, column=1 + j, value=t)

    # ---------- data sheets ----------
    for sp, (gl_row, pl_row) in (("NBA", (3, 4)), ("NFL", (5, 6))):
        cfg = SPORTS[sp]
        for kind, srow, fname in (("Data", gl_row, f"{sp.lower()}_gamelog.csv"),
                                  ("Players", pl_row, f"{sp.lower()}_players.csv")):
            ws = wb[f"{kind}_{sp}"]
            if mode == "template":
                ws["A1"] = f"=IMPORTDATA(Settings!$B${srow})"
            else:
                df = pd.read_csv(ROOT / "data" / fname)
                ws.append(list(df.columns))
                for row in df.itertuples(index=False):
                    ws.append([None if pd.isna(v) else (v.item() if hasattr(v, "item") else v) for v in row])

    for sp in ("NBA", "NFL"):
        dash(wb, sp)
        ticket(wb, sp)
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    print("saved", out)


def rng(sp, cfg, col):
    i = cfg["cols"].index(col) + 1
    return f"Data_{sp}!${L(i)}$2:${L(i)}${MAXR}"


def statrng(sp, cfg, stat_cell):
    last = L(len(cfg["cols"]))
    return (f"INDEX(Data_{sp}!$A$2:${last}${MAXR},0,MATCH({stat_cell},Data_{sp}!$A$1:${last}$1,0))")


def dash(wb, sp):
    cfg = SPORTS[sp]
    ws = wb[f"{sp}_Dash"]
    p, stat, line, opp, thr, price = cfg["default"][0], cfg["default"][1], cfg["default"][2], "", cfg["default"][4], cfg["default"][3]
    ws["A1"] = f"{sp} player lookup"
    ws["A1"].font = Font(bold=True, size=14)
    inputs = [("Player", p), ("Stat", stat), ("Line (X or more)", line), ("Opponent (optional)", None),
              ("Splits use seasons from", f"={cfg['cur']}-1"), ("Threshold for X rows", thr), ("Price to test (decimal)", price)]
    for i, (lab, val) in enumerate(inputs):
        r = 3 + i
        ws.cell(row=r, column=1, value=lab)
        c = ws.cell(row=r, column=2, value=val)
        c.fill = INPUT
    ws["A10"] = "Break even at that price"
    ws["B10"] = '=IF(B9="","",1/B9)'
    ws["B10"].number_format = "0.0%"
    dv = DataValidation(type="list", formula1=f"=Players_{sp}!$A$2:$A$3000", allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv); dv.add("B3")
    dv2 = DataValidation(type="list", formula1='"' + ",".join(cfg["stats"]) + '"', allow_blank=False)
    ws.add_data_validation(dv2); dv2.add("B4")
    col = "A" if sp == "NBA" else "B"
    dv3 = DataValidation(type="list", formula1=f"=Settings!${col}$15:${col}${14 + len(cfg['teams'])}", allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv3); dv3.add("B6")

    P = rng(sp, cfg, "player")
    SR = statrng(sp, cfg, "$B$4")
    UR = rng(sp, cfg, cfg["usage"])
    RK = rng(sp, cfg, "rk")
    SEAS = rng(sp, cfg, "season")
    top = 12
    hdr(ws, top, ["Split", "Games", "Avg stat", cfg["usage_label"], "Hit %", "Edge vs break even", "Flag"])
    rows = []
    rows.append(("This season", [(SEAS, f'{cfg["cur"]}')], False))
    for k in cfg["long"]:
        rows.append((f"Last {k}", [(RK, f'"<="&{k}')], False))
    rows.append(("Last season", [(SEAS, f'{cfg["cur"]}-1')], False))
    for name, spec in cfg["splits"]:
        base = [(SEAS, '">="&$B$7')]
        need_opp = False
        if spec == "OPP":
            base.append((rng(sp, cfg, "opp"), "$B$6")); need_opp = True
        elif spec == "USAGE":
            base.append((UR, '">="&$B$8'))
        else:
            for colname, crit in spec:
                base.append((rng(sp, cfg, colname), crit))
        label = name.replace("X", "the threshold") if spec == "USAGE" else name
        rows.append((label, base, need_opp))
    r = top + 1
    for name, crit, need_opp in rows:
        extra = "".join(f",{a},{b}" for a, b in crit)
        cnt = f"COUNTIFS({P},$B$3{extra})"
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=(f'=IF($B$6="","",{cnt})' if need_opp else f"={cnt}"))
        ws.cell(row=r, column=3, value=f'=IF(N(B{r})=0,"",AVERAGEIFS({SR},{P},$B$3{extra}))')
        ws.cell(row=r, column=4, value=f'=IF(N(B{r})=0,"",AVERAGEIFS({UR},{P},$B$3{extra}))')
        ws.cell(row=r, column=5, value=f'=IF(N(B{r})=0,"",COUNTIFS({P},$B$3{extra},{SR},">="&$B$5)/B{r})')
        ws.cell(row=r, column=6, value=f'=IF(OR(E{r}="",$B$10=""),"",E{r}-$B$10)')
        ws.cell(row=r, column=7, value=f'=IF(B{r}="","pick opponent",IF(B{r}=0,"no games",IF(B{r}<Settings!$B$9,"small sample","")))' if need_opp else f'=IF(B{r}=0,"no games",IF(B{r}<Settings!$B$9,"small sample",""))')
        ws.cell(row=r, column=3).number_format = "0.0"
        ws.cell(row=r, column=4).number_format = "0.0"
        ws.cell(row=r, column=5).number_format = "0%"
        ws.cell(row=r, column=6).number_format = "+0%;-0%"
        r += 1
    last = r - 1
    ws.conditional_formatting.add(f"F{top+1}:F{last}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="006100")))
    ws.conditional_formatting.add(f"F{top+1}:F{last}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="9C0006")))
    # last 10 table
    t = last + 3
    ws.cell(row=t - 1, column=1, value="Last 10 games").font = BOLD
    hdr(ws, t, ["Game", "Date", "Opp", "Home", cfg["usage"], "Stat", "Hit?"])
    KEY = rng(sp, cfg, "key")
    for k in range(1, 11):
        rr = t + k
        ws.cell(row=rr, column=1, value=k)
        m = f'MATCH($B$3&"|"&$A{rr},{KEY},0)'
        ws.cell(row=rr, column=2, value=f'=IFERROR(INDEX({rng(sp, cfg, "date")},{m}),"")')
        ws.cell(row=rr, column=3, value=f'=IFERROR(INDEX({rng(sp, cfg, "opp")},{m}),"")')
        ws.cell(row=rr, column=4, value=f'=IFERROR(IF(INDEX({rng(sp, cfg, "home")},{m})=1,"H","A"),"")')
        ws.cell(row=rr, column=5, value=f'=IFERROR(INDEX({UR},{m}),"")')
        ws.cell(row=rr, column=6, value=f'=IFERROR(INDEX({SR},{m}),"")')
        ws.cell(row=rr, column=7, value=f'=IF(F{rr}="","",IF(F{rr}>=$B$5,"yes","no"))')
    for c, w in zip("ABCDEFG", (30, 11, 10, 10, 9, 14, 14)):
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "A3"


def ticket(wb, sp):
    cfg = SPORTS[sp]
    ws = wb[f"{sp}_Ticket"]
    ws["A1"] = f"{sp} ticket check"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Fill the yellow cells. Leave 'My est' blank to use the auto estimate (L20 and season average, minus the haircut)."
    heads = ["Player", "Stat", "Line (X+)", "Price (decimal)", "Break even", "L10 hit", "L20 hit", "Season hit",
             "Season n", "Auto est", "My est", "Used est", "Edge", "Flags"]
    hdr(ws, 4, heads)
    P = rng(sp, cfg, "player"); RK = rng(sp, cfg, "rk"); SEAS = rng(sp, cfg, "season")
    first, n = 5, 20
    last = first + n - 1
    dv = DataValidation(type="list", formula1=f"=Players_{sp}!$A$2:$A$3000", allow_blank=True, showErrorMessage=False)
    dv2 = DataValidation(type="list", formula1='"' + ",".join(cfg["stats"]) + '"', allow_blank=True)
    ws.add_data_validation(dv); ws.add_data_validation(dv2)
    dv.add(f"A{first}:A{last}"); dv2.add(f"B{first}:B{last}")
    for i in range(n):
        r = first + i
        SR = statrng(sp, cfg, f"$B{r}")
        for c in "ABCDK":
            ws[f"{c}{r}"].fill = INPUT
        def hit(extra, nm):
            cnt = f"COUNTIFS({P},$A{r}{extra})"
            return f'=IF(OR($A{r}="",$B{r}="",$C{r}=""),"",IF({cnt}=0,"",COUNTIFS({P},$A{r}{extra},{SR},">="&$C{r})/{cnt}))'
        ws[f"E{r}"] = f'=IF(D{r}="","",1/D{r})'
        ws[f"F{r}"] = hit(f',{RK},"<=10"', "l10")
        ws[f"G{r}"] = hit(f',{RK},"<=20"', "l20")
        ws[f"H{r}"] = hit(f",{SEAS},{cfg['cur']}", "s")
        ws[f"I{r}"] = f'=IF($A{r}="","",COUNTIFS({P},$A{r},{SEAS},{cfg["cur"]}))'
        ws[f"J{r}"] = f'=IF(COUNT(G{r}:H{r})=0,"",AVERAGE(G{r}:H{r})*(1-Settings!$B$8))'
        ws[f"L{r}"] = f'=IF($A{r}="","",IF(K{r}<>"",K{r},J{r}))'
        ws[f"M{r}"] = f'=IF(OR(L{r}="",E{r}=""),"",L{r}-E{r})'
        ws[f"N{r}"] = (f'=IF($A{r}="","",TRIM(IF(COUNTIF($A${first}:$A${last},$A{r})>2,COUNTIF($A${first}:$A${last},$A{r})&" legs same player. ","")'
                       f'&IF(N(I{r})<Settings!$B$9,"small season sample. ","")&IF(AND(K{r}="",J{r}=""),"no data. ","")))')
        for c in "EFGHJKLM":
            ws[f"{c}{r}"].number_format = "0%" if c != "M" else "+0%;-0%"
        ws[f"E{r}"].number_format = "0.0%"
    ws.conditional_formatting.add(f"M{first}:M{last}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="006100")))
    ws.conditional_formatting.add(f"M{first}:M{last}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="9C0006")))
    s = last + 3
    ws.cell(row=s - 1, column=1, value="Ticket summary").font = BOLD
    summ = [
        ("Slip price before boost (decimal)", None, "0.00"),
        ("Profit boost %", 1.0, "0%"),
        ("Boosted decimal", f"=IF(B{s}=\"\",\"\",1+(B{s}-1)*(1+B{s+1}))", "0.00"),
        ("Break even after boost", f"=IF(B{s+2}=\"\",\"\",1/B{s+2})", "0.000%"),
        ("Legs entered", f'=COUNTA(A{first}:A{last})', "0"),
        ("Estimated hit chance (legs multiplied)", f'=IF(COUNT(L{first}:L{last})=0,"",PRODUCT(L{first}:L{last}))', "0.000%"),
        ("EV per $1 staked", f'=IF(OR(B{s+5}="",B{s+2}=""),"",B{s+5}*B{s+2}-1)', "+0.0%;-0.0%"),
        ("One in", f'=IF(N(B{s+5})=0,"",1/B{s+5})', "#,##0"),
        ("Stake", None, "0.00"),
        ("Payout if it hits", f'=IF(OR(B{s+8}="",B{s+2}=""),"",B{s+8}*B{s+2})', "#,##0.00"),
        ("Product of leg prices (compare to slip price)", f'=IF(COUNT(D{first}:D{last})=0,"",PRODUCT(D{first}:D{last}))', "0.00"),
    ]
    for i, (lab, val, fmt) in enumerate(summ):
        ws.cell(row=s + i, column=1, value=lab)
        c = ws.cell(row=s + i, column=2, value=val)
        c.number_format = fmt
        if val is None or lab.startswith("Profit"):
            c.fill = INPUT
    note = s + len(summ) + 1
    ws.cell(row=note, column=1, value="Note: legs multiplied as if independent. Same game legs are correlated and a book's slip price already accounts for that.")
    ws.cell(row=note + 1, column=1, value="If the product of leg prices is far above the slip price, the book is charging for correlation.")
    for c, w in zip("ABCDEFGHIJKLMN", (36, 11, 9, 9, 9, 8, 8, 9, 8, 8, 8, 8, 8, 34)):
        ws.column_dimensions[c].width = w
    ws.freeze_panes = "B5"


if __name__ == "__main__":
    mode = sys.argv[1] if len(sys.argv) > 1 else "template"
    name = "gamelog_TEST.xlsx" if mode == "test" else "gamelog_TEMPLATE.xlsx"
    build(mode, ROOT / "out" / name)
